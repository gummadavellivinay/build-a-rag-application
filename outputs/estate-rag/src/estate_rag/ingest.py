"""Offline ingestion boundary: files are processed by a trusted operator, never an API upload."""

import hashlib
import json
import re
from pathlib import Path

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter, TokenTextSplitter
from .models import Chunk


def ocr_image(image, languages: str) -> tuple[str, float]:
    import pytesseract
    from pytesseract import Output

    data = pytesseract.image_to_data(image, lang=languages, output_type=Output.DICT, timeout=90)
    words, scores = [], []
    for text, score in zip(data["text"], data["conf"]):
        if text.strip() and float(score) >= 0:
            words.append(text)
            scores.append(float(score))
    confidence = sum(scores) / len(scores) if scores else 0.0
    if confidence < 30 or not words:
        raise ValueError("OCR text is empty or low confidence; review the source before ingestion")
    return " ".join(words), confidence


def extract(path: Path, settings, force_ocr: bool = False) -> list[Document]:
    if path.stat().st_size > settings.max_file_mb * 1024 * 1024:
        raise ValueError(f"File too large: {path.name}")
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return [Document(page_content=path.read_text(encoding="utf-8"), metadata={"page": 1})]
    if suffix == ".docx":
        from docx import Document as WordDocument

        doc = WordDocument(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n" + "\n".join(" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows)
        return [Document(page_content=text, metadata={"page": 1})]
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp"}:
        from PIL import Image, ImageOps, ImageSequence

        result = []
        with Image.open(path) as image:
            for i, frame in enumerate(ImageSequence.Iterator(image)):
                if i >= settings.max_pages:
                    raise ValueError("Too many image frames")
                text, score = ocr_image(ImageOps.exif_transpose(frame).convert("RGB"), settings.ocr_languages)
                result.append(
                    Document(
                        page_content=text,
                        metadata={"page": i + 1, "modality": "ocr", "ocr_confidence": score},
                    )
                )
        return result
    if suffix == ".pdf":
        import pymupdf
        from PIL import Image

        result = []
        with pymupdf.open(path) as pdf:
            if pdf.is_encrypted or len(pdf) > settings.max_pages:
                raise ValueError("Encrypted PDF or page limit exceeded")
            for i, page in enumerate(pdf):
                native = page.get_text("text").strip()
                # OCR the entire rendered page if it contains raster images, even when a text layer exists.
                # This catches screenshots embedded in otherwise searchable PDFs.
                needs_ocr = force_ocr or len(native) < 40 or bool(page.get_images())
                if needs_ocr:
                    pix = page.get_pixmap(dpi=200, alpha=False, colorspace=pymupdf.csRGB)
                    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    text, score = ocr_image(image, settings.ocr_languages)
                    meta = {"page": i + 1, "modality": "ocr", "ocr_confidence": score}
                else:
                    text, meta = native, {"page": i + 1}
                if text:
                    result.append(Document(page_content=text, metadata=meta))
        return result
    raise ValueError(f"Unsupported format: {suffix}")


def split_documents(documents, strategy="recursive", embeddings=None):
    """400-token limit; semantic grouping is optional and incurs embedding calls."""
    recursive = RecursiveCharacterTextSplitter.from_tiktoken_encoder(
        encoding_name="cl100k_base",
        chunk_size=400,
        chunk_overlap=60,
        separators=["\n\n", "\n", ". ", "。", " ", ""],
    )
    if strategy == "fixed":
        return TokenTextSplitter(chunk_size=400, chunk_overlap=60).split_documents(documents)
    if strategy == "recursive":
        return recursive.split_documents(documents)
    if strategy != "semantic" or embeddings is None:
        raise ValueError("Semantic splitting requires an embeddings client")
    import numpy as np

    result = []
    for doc in documents:
        sentences = [s.strip() for s in re.split(r"(?<=[.!?。])\s+|\n\n", doc.page_content) if s.strip()]
        if not sentences:
            continue
        # Bound even single long sentences before embedding.
        units = [part for s in sentences for part in recursive.split_text(s)]
        vectors = np.asarray(embeddings.embed_documents(units), dtype=float)
        vectors /= np.maximum(np.linalg.norm(vectors, axis=1, keepdims=True), 1e-12)
        groups, group = [], []
        for i, unit in enumerate(units):
            if i and float(vectors[i] @ vectors[i - 1]) < 0.72:
                groups.append(" ".join(group))
                group = []
            group.append(unit)
        groups.append(" ".join(group))
        for text in groups:
            result.extend(recursive.split_documents([Document(page_content=text, metadata=doc.metadata)]))
    return result


def load_corpus(manifest_path: Path, tenant: str, settings, strategy="recursive", embeddings=None):
    records = json.loads(manifest_path.read_text())
    chunks = []
    for record in records:
        path = (manifest_path.parent / record["path"]).resolve()
        if not path.is_relative_to(manifest_path.parent.resolve()):
            raise ValueError("Manifest paths must remain inside the corpus directory")
        docs = extract(path, settings, record.get("force_ocr", False))
        for i, doc in enumerate(split_documents(docs, strategy, embeddings)):
            text = doc.page_content.replace("\x00", "").strip()
            if not text:
                continue
            identity = f"{tenant}:{record['path']}:{doc.metadata['page']}:{i}:{text}"
            chunks.append(
                Chunk(
                    id=hashlib.sha256(identity.encode()).hexdigest()[:32],
                    text=text,
                    source=record["path"],
                    tenant=tenant,
                    listing_id=record.get("listing_id", ""),
                    language=record.get("language", "und"),
                    **doc.metadata,
                )
            )
    if not chunks:
        raise ValueError("Corpus contains no indexable text")
    if len({c.id for c in chunks}) != len(chunks):
        raise ValueError("Duplicate records in corpus manifest")
    return chunks
